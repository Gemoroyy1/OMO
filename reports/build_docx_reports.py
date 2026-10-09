"""Редактируемые Word-отчёты из того же содержимого, что и PDF."""
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import Paragraph, Image, Table, Spacer, PageBreak, KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
import build_reports as source


def plain(value):
    return html.unescape(re.sub(r'<[^>]+>', '', value.text))


def field(paragraph, instruction, cached=''):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve')
    code.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for element in (begin, code, separate): run._r.append(element)
    if cached: paragraph.add_run(cached)
    paragraph.add_run()._r.append(end)


class WordReport:
    def __init__(self, path, **kwargs):
        self.output = Path(path).with_suffix('.docx')
        self.page_map = []

    def multiBuild(self, items):
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(30), Mm(15)
        section.top_margin = section.bottom_margin = Mm(20)
        section.footer_distance = Mm(10)
        section.different_first_page_header_footer = True
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
            style = doc.styles[name]
            style.font.name = 'Times New Roman'
            style.font.size = Pt(14)
            style.font.bold = False
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
        normal = doc.styles['Normal'].paragraph_format
        normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.first_line_indent = Mm(12.5)
        for name in ('Heading 1', 'Heading 2'):
            doc.styles[name].paragraph_format.keep_with_next = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(footer, ' PAGE ')
        update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true')
        doc.settings.element.append(update)
        doc.core_properties.title = self.output.parent.parent.name + ' Отчёт'
        doc.core_properties.author = ''
        headings = [plain(item) for item in items if isinstance(item, Paragraph) and hasattr(item, 'toc_title')]

        def append(item):
            if isinstance(item, KeepTogether):
                for child in item._content: append(child)
            elif isinstance(item, PageBreak):
                doc.add_page_break()
            elif isinstance(item, TableOfContents):
                # Cache readable entries; Word refreshes the field with actual page numbers.
                field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', '\n'.join(headings))
            elif isinstance(item, Paragraph):
                text = plain(item)
                style_name = item.style.name
                heading = hasattr(item, 'toc_title')
                p = doc.add_paragraph(text, style='Heading 1' if heading else ('Heading 2' if style_name == 'Heading' else 'Normal'))
                fmt = p.paragraph_format
                if style_name in ('Center', 'Caption'):
                    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fmt.first_line_indent = Mm(0)
                if style_name == 'TableTitle':
                    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    fmt.first_line_indent = Mm(0)
                    fmt.keep_with_next = True
                if style_name in ('Caption', 'TableTitle'):
                    for run in p.runs: run.font.size = Pt(12)
                if style_name == 'Code':
                    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    fmt.first_line_indent = Mm(0)
                    fmt.line_spacing = 1
                    fmt.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = 'Courier New'; run.font.size = Pt(12)
                if text == 'ОТЧЁТ': p.style = doc.styles['Title']
            elif isinstance(item, Image):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Mm(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(item.filename, width=Pt(item.drawWidth), height=Pt(item.drawHeight))
            elif isinstance(item, Table):
                rows = item._cellvalues
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                for column, width in zip(table.columns, item._colWidths): column.width = Pt(width)
                borders = OxmlElement('w:tblBorders')
                for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                    border = OxmlElement('w:' + edge)
                    for key, val in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items(): border.set(qn('w:' + key), val)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for r, values in enumerate(rows):
                    trpr = table.rows[r]._tr.get_or_add_trPr()
                    no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
                    if r == 0: trpr.append(OxmlElement('w:tblHeader'))
                    for c, value in enumerate(values):
                        cell = table.cell(r, c)
                        cell.width = Pt(item._colWidths[c])
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if r == 0:
                            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(fill)
                        p = cell.paragraphs[0]
                        p.paragraph_format.first_line_indent = Mm(0)
                        p.paragraph_format.line_spacing = 1
                        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run(plain(value).replace('train_RMSE', 'RMSE train'))
                        run.font.size = Pt(12)
            elif isinstance(item, Spacer):
                # Small gaps already supplied by paragraph styles; keep title-page spacing.
                if item.height >= 30:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(item.height)
                    p.paragraph_format.line_spacing = Pt(1)
        for item in items: append(item)
        doc.save(self.output)
        reopened = Document(self.output)
        assert reopened.tables and reopened.inline_shapes
        assert any('ЗАКЛЮЧЕНИЕ' in p.text for p in reopened.paragraphs)
        print(self.output)


if __name__ == '__main__':
    (source.ROOT / 'tmp').mkdir(exist_ok=True)
    source.Report = WordReport
    for number in (1, 2, 3): source.build(number)
